"""
Convert every .xls / .xlsx file found in temp/ (or another directory) to CSV.

Each worksheet becomes one CSV file:

    single-sheet workbook   ->  <name>.csv
    multi-sheet workbook    ->  <name>.<sheet>.csv   (one per sheet)

Sheet names are sanitised for use in filenames. Outputs that already exist on
disk are skipped unless --overwrite is given. If two workbooks in one run
would map to the same CSV name, the source extension is inserted to
disambiguate (e.g. foo.xls.csv). Excel lock files (~$*.xlsx) are ignored.

Usage:
    python scripts/convert_xls_to_csv.py
    python scripts/convert_xls_to_csv.py --dir temp --out temp/csv_out
    python scripts/convert_xls_to_csv.py --recursive --overwrite
    python scripts/convert_xls_to_csv.py --sheet "toata tara" --encoding utf-8-sig

Requirements:
    pip install openpyxl xlrd
    (openpyxl reads .xlsx; xlrd>=2 reads legacy .xls)
"""
import argparse
import csv
import re
import sys
from datetime import date, datetime, time
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE_DIR = PROJECT_ROOT / "temp"

XLS_SUFFIXES = {".xls", ".xlsx"}

# Characters not allowed in Windows filenames, plus control characters.
_UNSAFE_CHARS = re.compile(r'[\\/:*?"<>|\x00-\x1f]')

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import xlrd
except ImportError:
    xlrd = None


def sanitize_sheet_name(name: str) -> str:
    """Make a sheet name safe to use in a filename."""
    cleaned = _UNSAFE_CHARS.sub("_", str(name)).strip().strip(".")
    return cleaned[:100] or "sheet"


def format_value(value) -> object:
    """Normalise a cell value (openpyxl or xlrd flavour) for CSV output."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        # Date-formatted cells come back as datetime; emit plain dates
        # without a spurious midnight time component.
        if (value.hour, value.minute, value.second, value.microsecond) == (0, 0, 0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
        # Avoid "5.0" for integer-valued numeric cells (codes, counts, ...).
        return str(int(value))
    return value


def read_xlsx(path: Path) -> list[tuple[str, list[list]]]:
    """Read a .xlsx workbook, returning [(sheet_name, rows), ...]."""
    if openpyxl is None:
        raise RuntimeError(
            "openpyxl is required to read .xlsx files (pip install openpyxl)"
        )
    # data_only=True returns cached formula results instead of formula text.
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets: list[tuple[str, list[list]]] = []
        for sheet_name in workbook.sheetnames:
            rows = [
                [format_value(cell) for cell in row]
                for row in workbook[sheet_name].iter_rows(values_only=True)
            ]
            sheets.append((sheet_name, rows))
        return sheets
    finally:
        workbook.close()


def read_xls(path: Path) -> list[tuple[str, list[list]]]:
    """Read a legacy .xls workbook, returning [(sheet_name, rows), ...]."""
    if xlrd is None:
        raise RuntimeError(
            "xlrd is required to read legacy .xls files (pip install xlrd)"
        )
    book = xlrd.open_workbook(str(path))
    error_text = getattr(xlrd, "error_text_from_code", {})
    sheets: list[tuple[str, list[list]]] = []
    for sheet in book.sheets():
        rows: list[list] = []
        for r in range(sheet.nrows):
            row: list = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_TEXT:
                    row.append(cell.value)
                elif cell.ctype == xlrd.XL_CELL_NUMBER:
                    row.append(format_value(cell.value))
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    dt = xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
                    row.append(format_value(dt))
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    row.append("TRUE" if cell.value else "FALSE")
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append("")
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    row.append(error_text.get(cell.value, f"#ERR{cell.value}"))
                else:
                    row.append(str(cell.value))
            rows.append(row)
        sheets.append((sheet.name, rows))
    return sheets


def find_spreadsheet_files(directory: Path, recursive: bool) -> list[Path]:
    """Collect .xls/.xlsx files (skipping Excel lock files), sorted by name."""
    iterator = directory.rglob("*") if recursive else directory.iterdir()
    return sorted(
        path
        for path in iterator
        if path.is_file()
        and path.suffix.lower() in XLS_SUFFIXES
        and not path.name.startswith("~$")
    )


def csv_target(
    src: Path,
    sheet_name: str,
    multi_sheet: bool,
    out_dir: Path | None,
    used: set,
) -> Path:
    """Decide the output path for one sheet, avoiding in-run collisions."""
    base = out_dir if out_dir is not None else src.parent
    name = (
        f"{src.stem}.{sanitize_sheet_name(sheet_name)}.csv"
        if multi_sheet
        else f"{src.stem}.csv"
    )
    target = base / name
    if target in used:
        # e.g. foo.xls and foo.xlsx both present: keep both via the extension.
        parts = [src.stem, src.suffix.lstrip(".").lower()]
        if multi_sheet:
            parts.append(sanitize_sheet_name(sheet_name))
        target = base / (".".join(parts) + ".csv")
    return target


def convert_file(
    src: Path,
    out_dir: Path | None,
    overwrite: bool,
    encoding: str,
    sheet_filter: set[str],
    used: set,
) -> tuple[int, int]:
    """Convert one workbook; returns (csv_files_written, csv_files_skipped)."""
    reader = read_xlsx if src.suffix.lower() == ".xlsx" else read_xls
    sheets = reader(src)
    multi_sheet = len(sheets) > 1
    written = skipped = 0
    for sheet_name, rows in sheets:
        if sheet_filter and sheet_name.lower() not in sheet_filter:
            continue
        target = csv_target(src, sheet_name, multi_sheet, out_dir, used)
        if target in used:
            print(
                f"  [skip] {sheet_name!r}: could not determine a free output name",
                file=sys.stderr,
            )
            skipped += 1
            continue
        used.add(target)
        if target.exists() and not overwrite:
            print(f"  [skip] {target.name} already exists (use --overwrite)")
            skipped += 1
            continue
        target.parent.mkdir(parents=True, exist_ok=True)
        with open(target, "w", newline="", encoding=encoding) as handle:
            csv.writer(handle).writerows(rows)
        print(f"  [ok] {target.name} ({len(rows)} rows)")
        written += 1
    return written, skipped


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert all .xls/.xlsx files in a directory to CSV "
        "(one CSV per worksheet)."
    )
    parser.add_argument(
        "--dir",
        type=Path,
        default=DEFAULT_SOURCE_DIR,
        help=f"directory to scan (default: {DEFAULT_SOURCE_DIR})",
    )
    parser.add_argument(
        "--out",
        type=Path,
        default=None,
        help="output directory (default: next to each source file)",
    )
    parser.add_argument(
        "--recursive",
        action="store_true",
        help="also scan subdirectories of --dir",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="re-convert even if the output CSV already exists",
    )
    parser.add_argument(
        "--encoding",
        default="utf-8",
        help="output CSV encoding (default: utf-8; use utf-8-sig for Excel)",
    )
    parser.add_argument(
        "--sheet",
        action="append",
        default=None,
        metavar="NAME",
        help="convert only the named sheet(s); repeatable, case-insensitive",
    )
    args = parser.parse_args(argv)
    args.sheet_filter = {s.lower() for s in args.sheet} if args.sheet else set()
    return args


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    source_dir: Path = args.dir
    if not source_dir.is_dir():
        print(f"ERROR: directory not found: {source_dir}", file=sys.stderr)
        return 1

    files = find_spreadsheet_files(source_dir, args.recursive)
    if not files:
        print(f"No .xls/.xlsx files found in {source_dir}")
        return 0
    print(f"Found {len(files)} workbook(s) in {source_dir}\n")

    if args.out is not None:
        args.out.mkdir(parents=True, exist_ok=True)

    used: set = set()
    total_written = total_skipped = failures = 0
    for src in files:
        print(f"{src.name}:")
        try:
            written, skipped = convert_file(
                src, args.out, args.overwrite, args.encoding, args.sheet_filter, used
            )
        except Exception as exc:
            print(f"  [fail] {exc}", file=sys.stderr)
            failures += 1
            continue
        total_written += written
        total_skipped += skipped

    print(
        f"\nDone: {total_written} CSV file(s) written, "
        f"{total_skipped} skipped, {failures} workbook(s) failed."
    )
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())


